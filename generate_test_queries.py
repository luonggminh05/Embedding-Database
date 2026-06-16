import argparse
import re
import sys
from pathlib import Path
from typing import Iterable, List

import openpyxl
from docx import Document
from pypdf import PdfReader


QUESTION_TEMPLATES = [
    "{topic} là gì?",
    "Trình bày {topic}.",
    "Giải thích {topic}.",
    "Nêu các ý chính về {topic}.",
    "Phân tích {topic}.",
    "So sánh {topic} với các khái niệm liên quan.",
    "Ví dụ minh họa cho {topic} là gì?",
    "Tóm tắt nội dung về {topic}.",
    "Các đặc điểm của {topic} là gì?",
    "Vai trò của {topic} là gì?",
    "Vì sao cần hiểu {topic}?",
    "Ứng dụng của {topic} trong thực tế là gì?",
    "Các thành phần liên quan đến {topic} gồm những gì?",
    "Quy trình hoặc cách hoạt động của {topic} như thế nào?",
    "Những điểm cần lưu ý về {topic} là gì?",
    "Hãy đặt {topic} trong bối cảnh môn học.",
    "Câu hỏi ôn tập về {topic}.",
    "Những lỗi thường gặp khi hiểu {topic} là gì?",
    "Mối quan hệ giữa {topic} và nội dung trước đó là gì?",
    "Khi nào sử dụng {topic}?",
]

EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.\w+")
STUDENT_ID_RE = re.compile(r"\b\d{7,}\b")


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def read_pdf(path: Path) -> List[str]:
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = normalize_text(page.extract_text() or "")
        if text:
            pages.append(text)
    return pages


def read_docx(path: Path) -> List[str]:
    doc = Document(str(path))
    paragraphs = [normalize_text(p.text) for p in doc.paragraphs]
    return [paragraph for paragraph in paragraphs if paragraph]


def read_xlsx(path: Path) -> List[str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [normalize_text(str(value)) for value in row if value is not None]
            if values:
                rows.append(" - ".join(values))
    workbook.close()
    return rows


def read_supported_file(path: Path) -> List[str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    return []


def split_candidates(texts: Iterable[str]) -> List[str]:
    candidates = []
    for text in texts:
        parts = re.split(r"(?<=[.!?])\s+|\n+|•|·|[-–]\s+", text)
        for part in parts:
            part = normalize_text(part)
            if 12 <= len(part) <= 120:
                candidates.append(part)
    return candidates


def clean_topic(candidate: str) -> str:
    candidate = re.sub(r"^[\dIVXLCDMivxlcdm]+[\).:-]\s*", "", candidate)
    candidate = re.sub(r"^(chương|bài|câu|chapter)\s+\d+[\.: -]*", "", candidate, flags=re.IGNORECASE)
    candidate = candidate.strip(" .,:;!?")
    return normalize_text(candidate)


def looks_like_roster_row(text: str) -> bool:
    lowered = text.lower()
    if EMAIL_RE.search(text) or STUDENT_ID_RE.search(text):
        return True
    roster_words = ["mssv", "họ tên", "ho ten", "email", "địa chỉ thư điện tử"]
    return any(word in lowered for word in roster_words)


def is_good_topic(topic: str) -> bool:
    if looks_like_roster_row(topic):
        return False
    if len(topic) < 8 or len(topic.split()) > 18:
        return False
    letters = [char for char in topic if char.isalpha()]
    if len(letters) < 6:
        return False
    if len(letters) / max(1, len(topic)) < 0.45:
        return False
    return True


def build_queries(candidates: Iterable[str], limit: int) -> List[str]:
    seen_topics = set()
    queries = []
    template_index = 0

    for candidate in candidates:
        topic = clean_topic(candidate)
        key = topic.lower()
        if key in seen_topics or not is_good_topic(topic):
            continue

        seen_topics.add(key)
        for _ in range(4):
            template = QUESTION_TEMPLATES[template_index % len(QUESTION_TEMPLATES)]
            queries.append(template.format(topic=topic))
            template_index += 1
            if len(queries) >= limit:
                return queries

    return queries


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="Generate search queries from papers/")
    parser.add_argument("--papers-dir", default="papers")
    parser.add_argument("--output", default="generated_test_queries.txt")
    parser.add_argument("--limit", type=int, default=1000)
    args = parser.parse_args()

    papers_dir = Path(args.papers_dir)
    all_texts = []
    per_file_counts = {}

    paths = sorted(
        (path for path in papers_dir.iterdir() if path.is_file()),
        key=lambda path: 1 if path.suffix.lower() in {".xlsx", ".xlsm"} else 0,
    )

    for path in paths:
        texts = read_supported_file(path)
        if texts:
            per_file_counts[path.name] = len(texts)
            all_texts.extend(texts)

    candidates = split_candidates(all_texts)
    queries = build_queries(candidates, args.limit)

    output_path = Path(args.output)
    output_path.write_text("\n".join(queries) + "\n", encoding="utf-8")

    print(f"Read {len(per_file_counts)} files")
    for name, count in per_file_counts.items():
        print(f"- {name}: {count} text blocks")
    print(f"Generated {len(queries)} queries -> {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
